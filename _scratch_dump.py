import openpyxl, sys, glob, os

def cols(ws):
    out = {}
    for col in range(1, ws.max_column + 1):
        letter = openpyxl.utils.get_column_letter(col)
        out[letter] = ws.cell(row=1, column=col).value
    return out

def dump(path):
    wb = openpyxl.load_workbook(path, data_only=False)
    ws = wb[wb.sheetnames[0]]
    print("=" * 100)
    print("FILE:", os.path.basename(path))
    print("sheets:", wb.sheetnames, "max_row:", ws.max_row, "max_col:", ws.max_column)
    hdr = cols(ws)
    print("headers:", {k: v for k, v in hdr.items()})
    print("-" * 100)
    for r in range(3, ws.max_row + 1):
        vals = []
        for col in [1, 2, 3, 4, 6, 8, 10, 12, 15, 19, 20, 22, 23, 24, 25, 26, 27, 28, 29]:
            letter = openpyxl.utils.get_column_letter(col)
            c = ws.cell(row=r, column=col)
            v = c.value
            vals.append(f"{letter}={v!r}")
        print(f"r{r}: " + " | ".join(vals))
    wb.close()

for p in sys.argv[1:]:
    dump(p)
